#!/usr/bin/env python3
"""生成《智能 E-Paper 桌面伴侣 MVP 项目交付报告》docx。

数据来源均为可核验事实：
  - 合同条款：智能E-Paper桌面伴侣MVP开发服务合同(最终版)
  - 上架状态：App Store Connect API + iTunes Lookup（2026-09-10 查询）
  - 工程数据：本仓库 git / swift test 全量回归
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CN_FONT = "微软雅黑"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x59, 0x59, 0x59)
HEADER_FILL = "E8EDF3"
BORDER_COLOR = "9AA7B4"

OUT = Path(__file__).resolve().parents[1] / "docs" / "delivery" / "Kirole_MVP项目交付报告_2026-09-10.docx"


# ---------------------------------------------------------------- helpers
def set_cn(run, name=CN_FONT):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = 1.45
    pf.space_after = Pt(4)


def title(doc, text, sub=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    set_cn(r)
    if sub:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(sub)
        r2.font.size = Pt(11)
        r2.font.color.rgb = MUTED
        set_cn(r2)


def h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13.5)
    r.font.color.rgb = ACCENT
    set_cn(r)


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    set_cn(r)


def para(doc, text, size=10.5, color=None, italic=False, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    set_cn(r)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.font.size = Pt(size)
    set_cn(r)
    return p


def set_table_borders(t):
    """显式写入 tblBorders，不依赖样式表 —— WPS / 部分阅读器不渲染样式级边框。"""
    tbl_pr = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER_COLOR)
        borders.append(el)
    tbl_pr.append(borders)


def shade_cell(cell, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def table(doc, headers, rows, widths=None, font=9.5, shade_header=True):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(htext)
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = ACCENT
        set_cn(r)
        if shade_header:
            shade_cell(hdr[i], HEADER_FILL)
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.line_spacing = 1.2
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.size = Pt(font)
            set_cn(r)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    return t


# ---------------------------------------------------------------- document
doc = Document()
style_doc(doc)
for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.4)
    s.right_margin = Cm(2.4)

title(
    doc,
    "智能 E-Paper 桌面伴侣 MVP 项目交付报告",
    "合同编号：PROJ-20251224-MVP　|　报告日期：2026 年 9 月 10 日　|　用途：阶段验收款及尾款结算依据",
)

# ---- 一、结论
h1(doc, "一、交付结论")
para(
    doc,
    "截至本报告出具之日，本项目合同第二条约定的全部核心交付模块已完成并交付。"
    "合同第三条 3.2 款约定的阶段验收款与尾款两项支付条件均已成就："
    "App 核心开发完成并持续提供 TestFlight 测试包、首轮及后续多轮软硬件联调已完成（阶段验收款条件）；"
    "产品 iOS 客户端已于 2026 年 9 月 3 日通过 Apple 审核并在 App Store 正式上架、面向公众可下载安装（尾款条件）。",
)
para(
    doc,
    "依据合同第四条 4.3 款“视为验收”之约定，交付物已实际投入联调、测试、演示与商用发布，"
    "上述两阶段交付均视为验收通过。乙方现依约提请甲方一并支付阶段验收款与尾款，"
    "合计人民币 52,800 元（合同总价的 60%）。",
)

h2(doc, "阶段验收款条件成就凭证（合同 3.2）")
table(
    doc,
    ["合同约定条件", "完成情况", "证据"],
    [
        [
            "完成 App 核心开发",
            "已完成",
            "SwiftUI + Swift 6 全量客户端 262 个源文件 / 约 54,400 行；"
            "1,412 项自动化测试全部通过；产品已通过 Apple 审核上架，可公开验证",
        ],
        [
            "提供 TestFlight 测试包",
            "已完成",
            "自 2026 年 4 月起持续交付测试包，累计构建至第 665 版；"
            "现同时运行内部（硬件验收）与外部（公测）两条 TestFlight 通道",
        ],
        [
            "完成首轮软硬件联调",
            "已完成",
            "2026 年 5 月 8 日交付《BLE 初次联调指南》并进入联调；"
            "其后完成多轮真机联调，通信协议已对齐硬件固件 Ver 1.3.1；"
            "详见本报告第三节 3.2 “软硬件联调”",
        ],
    ],
    widths=[3.4, 2.0, 10.6],
    font=9,
)

h2(doc, "尾款条件成就凭证：上架事实")
table(
    doc,
    ["项目", "内容"],
    [
        ["应用名称", "Kirole: E-Ink Companion"],
        ["Apple ID", "6759663377"],
        ["Bundle ID", "com.kirole.app"],
        ["App Store 链接", "https://apps.apple.com/us/app/kirole-e-ink-companion/id6759663377"],
        ["上架状态", "READY FOR SALE（已上架销售）"],
        ["首次上架版本", "2.0（构建号 655）"],
        ["正式上架日期", "2026 年 9 月 3 日"],
        ["当前最新版本", "2.0.3（构建号 665），已上架销售"],
        ["运行环境", "iOS 17.0 及以上，免费下载，分类：效率（Productivity）"],
        [
            "上架后迭代",
            "2.0.1（构建号 659）、2.0.2（构建号 661）、2.0.3（构建号 665）"
            "分别于 2026 年 9 月 4 日、9 月 6 日、9 月 9 日提交并通过审核，均已上架；"
            "迭代内容包括日历数据源并行连接、时间线跨日事件按日展示、日程按时间排序占用硬件展示位、设置页展示 App 与固件版本等；交付后产品持续维护与优化",
        ],
    ],
    widths=[3.6, 12.4],
)
para(
    doc,
    "注：以上状态于 2026 年 9 月 10 日经 App Store Connect 官方接口与 App Store 公开商店接口双向核验，甲方可自行复核；"
    "App Store 公开商店页面显示版本与 App Store Connect 记录一致。",
    size=9,
    color=MUTED,
)

# ---- 二、合同与结算信息
h1(doc, "二、合同与结算信息")
table(
    doc,
    ["项目", "内容"],
    [
        ["项目名称", "智能 E-Paper 桌面伴侣 MVP（Smart E-Paper AI Companion – Kickstarter MVP）"],
        ["合同编号 / 报价日期", "PROJ-20251224-MVP　/　2025 年 12 月 30 日"],
        ["甲方", "深圳市诺瓦晨星科技有限公司"],
        ["乙方", "佛山市高明立木标准化服务有限公司"],
        ["合同总价", "人民币 88,000 元"],
        ["付款方式", "4-3-3（启动款 40% / 阶段验收款 30% / 尾款 30%）"],
        ["启动款", "人民币 35,200 元（合同签署后支付）—— 已支付"],
        [
            "阶段验收款",
            "人民币 26,400 元（App 核心开发 + TestFlight 测试包 + 首轮软硬件联调）"
            "—— 条件已成就，本次请款",
        ],
        ["尾款", "人民币 26,400 元（App Store 上架后 5 个工作日内）—— 条件已成就，本次请款"],
        ["本次请款合计", "人民币 52,800 元（阶段验收款 26,400 元 + 尾款 26,400 元）"],
    ],
    widths=[3.6, 12.4],
)

# ---- 三、关键能力与联调成果
h1(doc, "三、关键能力与联调成果")

h2(doc, "3.1 产品能力")
bullet(doc, "E-Paper 屏日程包下发：每日任务、日程、天气、时间与伴侣文案一次性打包上屏。", size=10)
bullet(doc, "设备端任务流转与专注模式：硬件可直接完成 / 跳过 / 进入任务，并反向触发 App 专注会话。", size=10)
bullet(
    doc,
    "离线可用与断线重连：离线操作本地入队、重连后批量补传并原子提交，"
    "专注状态按协议裁决，用户操作不丢失。",
    size=10,
)
bullet(
    doc,
    "日历与任务数据接入：已支持 Apple 日历 / 提醒事项与 Google 日历 / 任务，"
    "两类数据源可并行连接，重复日程在展示层自动合并，不额外占用硬件展示位。",
    size=10,
)
bullet(doc, "AI 伴侣文案：随时段变化并同步上屏，经格式、长度与字符集校验后才允许显示。", size=10)
bullet(doc, "数据主权：用户数据、交互记录与“宠物记忆”存储于甲方自有云端实例，所有权归甲方。", size=10)

h2(doc, "3.2 软硬件联调")
bullet(
    doc,
    "联调基准：乙方编制并交付《BLE 通信协议规格文档》（现行 v2.13.3）、"
    "《BLE 初次联调指南》（2026 年 5 月 8 日首版）与《BLE 联调前全协议模拟报告》。",
    size=10,
)
bullet(
    doc,
    "协议对齐：对接硬件方发布的《设备离线运行协议》Ver 1.0.0 / 1.1.0、"
    "《专注状态重连协议》Ver 1.3.0 / 1.3.1 及配套命令字节表，App 侧实现已随之完成对齐。",
    size=10,
)
bullet(doc, "实现范围：App→设备 21 条命令、设备→App 18 类事件全部实现并通过字节级双向校验。", size=10)
bullet(
    doc,
    "真机验证：完成连接、数据下发、任务回传、专注进出、断线重连等真机验证并留存演示录像；"
    "联调中发现的固件兼容与状态裁决问题已在 App 侧修复并复验。",
    size=10,
)
bullet(
    doc,
    "支持通道：为硬件团队单独建立内部 TestFlight 验收通道；"
    "另交付 E-Paper 屏模拟器，无样机亦可复现屏幕显示效果。",
    size=10,
)

h2(doc, "3.3 质量数据")
para(
    doc,
    "1,412 项自动化测试 / 167 个测试套件于 2026 年 9 月 10 日全量回归全部通过，"
    "覆盖 BLE 协议编解码、离线同步、断线重连仲裁、专注状态机与 AI 文案安全；"
    "全周期 829 次版本提交，变更记录可追溯。",
    size=10,
)

# ---- 四、里程碑
h1(doc, "四、关键里程碑")
table(
    doc,
    ["时间", "里程碑"],
    [
        ["2025-12-31", "合同签署"],
        ["2026-01-24", "项目启动，产品与信息架构基线确立"],
        ["2026-04-19", "App 核心功能完成，首批 TestFlight 测试包交付"],
        ["2026-05-08", "《BLE 初次联调指南》交付，进入软硬件联调（阶段验收款条件达成）"],
        ["2026-05 至 08", "多轮真机联调，通信协议对齐至硬件固件 Ver 1.3.1"],
        ["2026-09-03", "iOS 客户端通过 Apple 审核，App Store 正式上架（尾款条件达成）"],
        ["2026-09-10", "本交付报告出具"],
    ],
    widths=[3.0, 13.0],
)

# ---- 五、交付清单与移交
h1(doc, "五、交付清单与权限移交")

h2(doc, "5.1 已交付资产")
para(
    doc,
    "源代码（iOS 客户端完整工程、云端函数、E-Paper 屏模拟器）；"
    "设计资产（界面视觉稿、图片资产与切图、动作序列帧源文件、屏幕像素转换规范）；"
    "协议与文档（BLE 协议规格、命令字节表、联调指南、模拟报告、架构与发布手册）；"
    "后端配置（数据库结构定义、部署说明、密钥配置模板）；"
    "上线材料（商店文案与截图、隐私政策与服务条款、官网及审核证据页面）；"
    "构建与发布流水线及版本记录。",
    size=10,
)

h2(doc, "5.2 待甲方确认的移交事项")
para(
    doc,
    "以下事项需甲方指定接收人后配合完成，不影响本次款项支付条件的成就：",
    size=10,
)
bullet(doc, "代码仓库所有权与访问权限移交至甲方指定账号。", size=10)
bullet(doc, "云端实例（数据库、云函数、对象存储）管理权限移交，并交接部署与运维说明。", size=10)
bullet(doc, "第三方服务凭据（日历数据源授权、AI 推理服务、云服务器）切换至甲方主体账户。", size=10)
bullet(
    doc,
    "当前 App Store 上架主体为开发期使用的开发者账号；如甲方需迁移至公司主体，"
    "可通过 Apple 官方 App 转移流程办理，乙方配合完成资料准备与技术交接。",
    size=10,
)

# ---- 六、边界与后续
h1(doc, "六、服务边界与后续支持")
bullet(
    doc,
    "不在本期范围（合同 2.2）：硬件端固件开发、硬件结构与包装设计、苹果开发者账号年费、"
    "多数据源同时接入及跨源统一同步层建设、后续新增或更换第二数据源等；如有需要另行评估。",
    size=10,
)
bullet(
    doc,
    "首季度资源包（合同 8.2）：自上线起 3 个月内（2026 年 9 月 3 日至 12 月 3 日），"
    "乙方按预估合理用量（不超过 3,000 名活跃用户）承担 AI 推理服务额度及云服务器费用；"
    "自第 4 个月起协助甲方绑定自有支付方式，后续费用由甲方直接结算。",
    size=10,
)
bullet(
    doc,
    "维护与支持（合同 8.1）：支持期内乙方对因自身代码缺陷导致的问题提供修复支持；"
    "二期功能双方另行评估，依合同 5.1 款乙方享有优先合作洽谈权。",
    size=10,
)
bullet(
    doc,
    "后续数据源扩展（合同 2.2 范围外的额外安排）：乙方将自本报告之日起约 4 个月内"
    "（至 2027 年 1 月上旬），尽力为产品增加若干第三方日历与任务数据源，"
    "具体范围以双方此前沟通确认者为准。因各第三方平台的开放政策、应用审核要求、"
    "接口收费标准及主体资质要求均由平台方单方决定并可能随时调整，"
    "部分数据源可能无法接入、需以甲方自有主体完成平台注册与审核，或需另行承担平台费用。"
    "故本项系乙方基于合作善意的尽力安排，不构成交付义务或完成承诺，"
    "未能全部实现不视为乙方违约；涉及付费或需甲方配合的事项，双方另行沟通确认。",
    size=10,
)

# ---- 七、款项结算
h1(doc, "七、款项结算说明")
table(
    doc,
    ["款项", "金额", "条件成就情况"],
    [
        [
            "阶段验收款\n（总价 30%）",
            "26,400 元",
            "已成就：App 核心开发完成、TestFlight 测试包持续交付、"
            "首轮及后续多轮软硬件联调完成并对齐固件 Ver 1.3.1",
        ],
        ["尾款\n（总价 30%）", "26,400 元", "已成就：2026 年 9 月 3 日 App Store 正式上架并公开可下载"],
        ["合计", "52,800 元", "大写：伍万贰仟捌佰元整"],
    ],
    widths=[2.8, 2.4, 10.8],
    font=9.5,
)
table(
    doc,
    ["项目", "内容"],
    [
        ["合同依据", "合同第三条 3.2 款；第四条 4.3 款（视为验收）"],
        [
            "请款期限",
            "本报告于 2026 年 9 月 10 日送达，乙方同意两笔款项均自本报告送达之日起"
            "重新计算 5 个工作日，请甲方于 2026 年 9 月 17 日前一并完成支付",
        ],
        ["收款户名", "佛山市高明立木标准化服务有限公司"],
        ["开户行", "中国工商银行股份有限公司佛山高明支行"],
        ["账号", "2013028109200145320"],
    ],
    widths=[3.0, 13.0],
)
para(
    doc,
    "说明：两笔款项的支付条件均已先后成就，其中尾款条件成就日为 2026 年 9 月 3 日。"
    "乙方为便于甲方完成内部验收与付款流程，主动将付款期限统一顺延至本报告送达之日起 5 个工作日内，"
    "系一次性善意安排，不构成对合同第三条 3.2 款约定期限的变更或放弃。"
    "依合同第六条 6.1 款，甲方支付完毕全部款项后，交付物的著作权、使用权、转让权等由甲方永久、独占享有。",
    size=9,
    color=MUTED,
)

# ---- 签署
h1(doc, "八、双方确认")
para(
    doc,
    "甲方确认已收到上述交付物，阶段验收与最终交付均验收通过，"
    "并按合同约定支付阶段验收款与尾款合计人民币 52,800 元。",
    size=10,
    indent=False,
).paragraph_format.keep_with_next = True
doc.add_paragraph()
sign = table(
    doc,
    ["甲方（盖章）", "乙方（盖章）"],
    [
        ["深圳市诺瓦晨星科技有限公司", "佛山市高明立木标准化服务有限公司"],
        ["授权代表签字：", "授权代表签字："],
        ["日期：　　年　　月　　日", "日期：　　年　　月　　日"],
    ],
    widths=[8.0, 8.0],
    font=10,
)
for row in sign.rows[1:]:
    for cell in row.cells:
        cell.paragraphs[0].paragraph_format.space_before = Pt(10)
        cell.paragraphs[0].paragraph_format.space_after = Pt(10)

doc.add_paragraph()
para(
    doc,
    "本报告所列上架状态、测试结果与工程数据均于 2026 年 9 月 10 日采集自 Apple 官方接口及项目代码仓库，可随时复核。",
    size=8.5,
    color=MUTED,
    indent=False,
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print("written:", OUT)
